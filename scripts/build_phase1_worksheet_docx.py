from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(55, 65, 81)
        else:
            paragraph.add_run(part)


def clean_table_cell(text: str) -> str:
    return text.strip().replace("<br>", "\n")


def build(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for name, size, color in (
        ("Title", 23, RGBColor(15, 23, 42)),
        ("Heading 1", 16, RGBColor(8, 145, 178)),
        ("Heading 2", 13, RGBColor(15, 23, 42)),
        ("Heading 3", 11, RGBColor(234, 88, 12)),
    ):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color

    header = section.header.paragraphs[0]
    header.text = "CircleGuard · Worksheet 5 · Phase 1"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("BMad 6.12.0 · WDS v0.4.3 · 2026-09-14")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Consolas"
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(30, 41, 59)
                in_code = False
                code_lines = []
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[i + 1]):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[i]):
                    rows.append([clean_table_cell(c) for c in lines[i].strip("|").split("|")])
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.autofit = True
            for r_idx, row in enumerate(rows):
                for c_idx, value in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    set_cell_margins(cell)
                    if r_idx == 0:
                        set_cell_shading(cell, "0F172A")
                    p = cell.paragraphs[0]
                    add_inline(p, value)
                    for run in p.runs:
                        run.font.size = Pt(8 if cols >= 4 else 9)
                        if r_idx == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if not line.strip():
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p = doc.add_heading(level=level - 1)
            add_inline(p, text)
            i += 1
            continue
        bullet = re.match(r"^\s*-\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or numbered:
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(p, (bullet or numbered).group(1))
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.core_properties.title = "Worksheet 5 — Phase 1: Automated Audit & Benchmarking"
    doc.core_properties.subject = "CircleGuard BMad 6.10 to 6.12 audit benchmark"
    doc.core_properties.author = "CircleGuard Triad"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_phase1_worksheet_docx.py SOURCE.md OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
